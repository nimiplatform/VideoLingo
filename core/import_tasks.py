"""Original URL downloads and Excel task columns; no model execution."""
from pathlib import Path
from urllib.parse import urlparse
from zipfile import ZipFile


def zip_subtitles(request, workspace):
    from zipfile import ZIP_DEFLATED
    names = {'src.srt', 'trans.srt', 'src_trans.srt', 'trans_src.srt', 'audio/src_subs_for_audio.srt', 'audio/trans_subs_for_audio.srt'}
    output = workspace / 'subtitles.zip'
    with ZipFile(output, 'w', compression=ZIP_DEFLATED) as archive:
        for item in request['files']:
            source = Path(item['file']).resolve(strict=True)
            source.relative_to(workspace)
            if item['name'] not in names:
                raise ValueError('Unexpected subtitle archive entry')
            archive.write(source, item['name'])
    return {'file': str(output)}


def read_task_sheet(request, workspace):
    from openpyxl import load_workbook

    filename = Path(request['file']).resolve(strict=True)
    filename.relative_to(workspace)
    with ZipFile(filename) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 32 * 1024 ** 2:
            raise ValueError('The workbook expands beyond 32 MiB')
    book = load_workbook(filename, read_only=True, data_only=False, keep_links=False)
    columns = ['Video File', 'Source Language', 'Target Language', 'Dubbing', 'Status']
    try:
        sheet = book.active
        if sheet.max_row > 1001 or sheet.max_column > 32:
            raise ValueError('Use a task sheet with at most 1000 rows and 32 columns')
        cells = sheet.iter_rows()
        header = [str(cell.value).strip() if cell.value is not None else '' for cell in next(cells)]
        if any(header.count(name) != 1 for name in columns):
            raise ValueError('Expected one of each column: ' + ', '.join(columns))
        result = []
        for number, row in enumerate(cells, 2):
            if not any(cell.value is not None for cell in row):
                continue
            if any(cell.data_type == 'f' for cell in row):
                raise ValueError(f'Row {number}: use literal task values, not formulas')
            values = {name: row[header.index(name)].value for name in columns}
            source = str(values['Video File'] or '').strip()
            language = str(values['Source Language'] or '').strip()
            language_specified = bool(language)
            if language.lower() == 'auto':
                language = ''
            target = str(values['Target Language'] or '').strip()
            status = str(values['Status'] or '').strip()
            dub = values['Dubbing']
            if not source or len(source) > 4096:
                raise ValueError(f'Row {number}: missing or oversized Video File')
            if language not in {'', 'en', 'zh', 'ja', 'fr', 'de', 'it', 'ru', 'es'} or len(target) > 80:
                raise ValueError(f'Row {number}: unsupported source language or target name')
            if dub not in (None, '', 0, 1, '0', '1'):
                raise ValueError(f'Row {number}: Dubbing must be 0 or 1')
            is_url = urlparse(source).scheme.lower() in {'http', 'https'}
            if is_url and not urlparse(source).netloc:
                raise ValueError(f'Row {number}: invalid media URL')
            if not is_url and (Path(source).name != source or '/' in source or '\\' in source):
                raise ValueError(f'Row {number}: select the local media by filename, without directory paths')
            result.append({'row': number, 'videoFile': source, 'isUrl': is_url, 'sourceLanguage': language, 'sourceLanguageSpecified': language_specified,
                           'targetLanguage': target, 'dubbing': dub in (1, '1'), 'status': status,
                           'selected': not status or 'error' in status.lower()})
        if not result:
            raise ValueError('The task sheet contains no media rows')
        return {'rows': result}
    finally:
        book.close()


def download_url(request, workspace):
    from ruamel.yaml import YAML
    from core._1_ytdlp import download_video_ytdlp, find_video_files

    url = request['url']
    parsed = urlparse(url)
    if parsed.scheme not in {'https', 'http'} or not parsed.netloc:
        raise ValueError('Use an HTTP or HTTPS media URL')
    resolution = request['resolution']
    if resolution not in {'360', '480', '720', '1080', 'best'}:
        raise ValueError('Choose a supported download resolution')
    cookies = request.get('cookiesPath', '')
    if cookies:
        Path(cookies).resolve(strict=True).relative_to(workspace)
    youtube = {'cookies_path': cookies}
    if request.get('proxy') is not None:
        youtube['proxy'] = request['proxy']
    with open('config.yaml', 'w', encoding='utf-8') as file:
        YAML().dump({'youtube': youtube, 'allowed_video_formats': ['mp4', 'mov', 'avi', 'mkv', 'flv', 'wmv', 'webm']}, file)
    download_video_ytdlp(url, resolution=resolution)
    video = Path(find_video_files()).resolve(strict=True)
    video.relative_to(workspace)
    if not (0 < video.stat().st_size <= 4 * 1024 ** 3):
        raise ValueError('Downloaded media must be non-empty and at most 4 GB')
    mime = {'mp4': 'video/mp4', 'mov': 'video/quicktime', 'avi': 'video/x-msvideo', 'mkv': 'video/x-matroska',
            'flv': 'video/x-flv', 'wmv': 'video/x-ms-wmv', 'webm': 'video/webm'}[video.suffix[1:].lower()]
    return {'file': str(video), 'name': video.name, 'mediaType': mime}
